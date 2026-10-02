#!/usr/bin/env python3
"""Rebuild the DriveClarify deck with one Image2 background per slide.

The generated images are visual treatment only. Exact text, tables, paper figures,
and repository screenshots remain deterministic PowerPoint foreground objects.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
V1 = ROOT / "deliverables" / "driveclarify_research_progress_v1"
OUT = ROOT / "deliverables" / "driveclarify_research_progress_v2_image2"
BACKGROUNDS = OUT / "backgrounds"
SOURCE_PPTX = V1 / "DriveClarify_研究进展汇报_V1_中文.pptx"
EDITABLE_PPTX = OUT / "DriveClarify_研究进展汇报_V2_Image2_可编辑版.pptx"
FLAT_PPTX = OUT / "DriveClarify_研究进展汇报_V2_Image2_全页图像版.pptx"


REAL_ASSETS = {
    "SimLingo_CVPR2025_Fig2": V1 / "figures" / "paper_simlingo_fig2.png",
    "LMDrive_CVPR2024_Fig4": V1 / "figures" / "paper_lmdrive_fig4.png",
    "DriveGPT4_2023_Fig2": V1 / "figures" / "paper_drivegpt4_fig2.png",
    "Talk2Car_2019_Fig1": V1 / "figures" / "paper_talk2car_fig1.png",
    "KnowNo_2023_Fig1": V1 / "figures" / "paper_knowno_fig1.png",
    "CARLA_ACT_CONTROL_TICK_ACTIVE": V1 / "figures" / "repo_closed_loop_act_tick.png",
}

GLASS_COLORS = {
    "FFFFFF",
    "F6F8FB",
    "EEF2F7",
    "E6F7F7",
    "E8F7FA",
    "FFF4D8",
    "E9F6F1",
    "FCEBEC",
    "F0ECFB",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def move_shape_to_index(slide, shape, index: int) -> None:
    element = shape._element
    tree = slide.shapes._spTree
    tree.remove(element)
    tree.insert(index, element)


def add_picture_cover(slide, path: Path, slide_w: int, slide_h: int):
    with Image.open(path) as image:
        source_ratio = image.width / image.height
    frame_ratio = slide_w / slide_h
    picture = slide.shapes.add_picture(str(path), 0, 0, width=slide_w, height=slide_h)
    if source_ratio > frame_ratio:
        visible = frame_ratio / source_ratio
        crop = (1.0 - visible) / 2.0
        picture.crop_left = crop
        picture.crop_right = crop
    else:
        visible = source_ratio / frame_ratio
        crop = (1.0 - visible) / 2.0
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def set_shape_alpha(shape, opacity_percent: int) -> None:
    """Set solid-fill opacity using DrawingML alpha (100000 = opaque)."""
    solid_fill = shape._element.spPr.solidFill
    if solid_fill is None:
        return
    color = solid_fill.find("{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
    if color is None:
        return
    for old in list(color.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}alpha")):
        color.remove(old)
    alpha = OxmlElement("a:alpha")
    alpha.set("val", str(max(0, min(100, opacity_percent)) * 1000))
    color.append(alpha)


def remove_replaced_illustrations(slide, slide_number: int, slide_w: int, slide_h: int) -> None:
    """Remove only prior ImageGen illustrations that the new Image2 background replaces."""
    for shape in list(slide.shapes):
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        if slide_number == 1:
            if shape.left <= Inches(0.05) and shape.top <= Inches(0.05) and shape.width >= slide_w * 0.95 and shape.height >= slide_h * 0.95:
                slide.shapes._spTree.remove(shape._element)
        elif slide_number == 4:
            if shape.left < Inches(1.0) and Inches(1.0) < shape.top < Inches(2.0) and shape.width > Inches(6.0):
                slide.shapes._spTree.remove(shape._element)


def add_title_veil(slide, slide_w: int) -> None:
    veil = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, slide_w, Inches(1.22))
    veil.fill.solid()
    veil.fill.fore_color.rgb = RGBColor(248, 250, 252)
    veil.line.fill.background()
    set_shape_alpha(veil, 90)
    move_shape_to_index(slide, veil, 3)


def apply_glass_effect(slide) -> None:
    """Make pale foreground panels slightly translucent over Image2 artwork."""
    for shape in slide.shapes:
        if getattr(shape, "has_table", False) or shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            continue
        try:
            if shape.fill.type != MSO_FILL_TYPE.SOLID:
                continue
            value = str(shape.fill.fore_color.rgb)
        except (AttributeError, TypeError, ValueError):
            continue
        if value in GLASS_COLORS:
            set_shape_alpha(shape, 91 if value == "FFFFFF" else 94)


def fix_known_layout_issues(slide, slide_number: int) -> None:
    """Repair two source text boxes whose stored heights were zero/negative."""
    if slide_number == 7:
        for shape in slide.shapes:
            text = getattr(shape, "text", "").strip()
            if text.startswith("uncertainty estimation"):
                shape.top = Inches(5.96)
                shape.width = Inches(4.65)
                shape.height = Inches(0.48)
            elif text == "核心区别":
                shape.top = Inches(5.55)
            elif (
                shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                and Inches(7.0) <= shape.left <= Inches(7.1)
                and Inches(5.4) <= shape.top <= Inches(5.55)
            ):
                shape.height = Inches(1.02)

    if slide_number == 18:
        for shape in slide.shapes:
            text = getattr(shape, "text", "").strip()
            if text.startswith("scenario annotation"):
                shape.top = Inches(6.10)
                shape.width = Inches(6.30)
                shape.height = Inches(0.40)
            elif text == "annotation ≠ runtime output":
                shape.top = Inches(5.75)
            elif (
                shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                and Inches(5.3) <= shape.left <= Inches(5.45)
                and Inches(5.6) <= shape.top <= Inches(5.7)
            ):
                shape.height = Inches(0.91)


def copy_support_files() -> None:
    for name in [
        "DriveClarify_研究进展汇报_V1_逐页讲稿.md",
        "claim_evidence_matrix.csv",
        "literature_comparison.csv",
        "citation_list.md",
        "figure_sources.md",
    ]:
        source = V1 / name
        if source.exists():
            shutil.copy2(source, OUT / name.replace("V1", "V2_Image2"))


def write_image2_provenance() -> Path:
    manifest_path = OUT / "image2_prompt_manifest.json"
    prompts = json.loads(manifest_path.read_text(encoding="utf-8"))
    lines = [
        "# DriveClarify V2 Image2 生成溯源",
        "",
        "- 模式：内置 ImageGen（gpt-image-2 工作流），逐页独立生成。",
        "- 用途分类：scientific-educational / productivity-visual。",
        "- 内容策略：Image2 只生成视觉底图；中文正文、数字、表格、论文原图和 CARLA 截图均采用确定性前景叠加。",
        "- 共同约束：16:9 学术汇报构图；无文字、无数字、无 Logo、无水印、无虚构实验结果。",
        "",
    ]
    for item in prompts:
        number = item["slide"]
        lines.extend(
            [
                f"## Slide {number:02d} · {item['title']}",
                "",
                f"- 工作区文件：`backgrounds/slide_{number:02d}.png`",
                f"- 保留策略：{item['preserve']}",
                f"- 最终提示词：{item['prompt']}",
                "",
            ]
        )
    output = OUT / "image2_provenance.md"
    output.write_text("\n".join(lines), encoding="utf-8")

    figure_sources = OUT / "figure_sources.md"
    if figure_sources.exists():
        with figure_sources.open("a", encoding="utf-8") as handle:
            handle.write(
                "\n\n## V2 Image2 逐页底图\n\n"
                "第 1–25 页均使用独立 Image2 背景；完整逐页提示词、工作区路径与保留策略见 "
                "`image2_provenance.md`。这些背景不是实验结果。论文原图和 CARLA 截图仍为原始字节。\n"
            )
    return output


def write_background_manifest() -> Path:
    prompt_items = {
        item["slide"]: item
        for item in json.loads((OUT / "image2_prompt_manifest.json").read_text(encoding="utf-8"))
    }
    records = []
    for index in range(1, 26):
        path = BACKGROUNDS / f"slide_{index:02d}.png"
        with Image.open(path) as image:
            width, height = image.size
            mode = image.mode
        records.append(
            {
                "slide": index,
                "title": prompt_items[index]["title"],
                "path": str(path),
                "width": width,
                "height": height,
                "mode": mode,
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
                "generator": "built-in ImageGen / gpt-image-2 workflow",
            }
        )
    output = OUT / "background_manifest.json"
    output.write_text(json.dumps({"count": 25, "backgrounds": records}, ensure_ascii=False, indent=2), encoding="utf-8")
    return output


def compose() -> Path:
    missing = [BACKGROUNDS / f"slide_{index:02d}.png" for index in range(1, 26) if not (BACKGROUNDS / f"slide_{index:02d}.png").exists()]
    if missing:
        raise FileNotFoundError("Missing Image2 backgrounds:\n" + "\n".join(str(path) for path in missing))

    prs = Presentation(SOURCE_PPTX)
    if len(prs.slides) != 25:
        raise ValueError(f"Expected 25 source slides, found {len(prs.slides)}")

    for index, slide in enumerate(prs.slides, 1):
        fix_known_layout_issues(slide, index)
        remove_replaced_illustrations(slide, index, prs.slide_width, prs.slide_height)
        if index != 1:
            apply_glass_effect(slide)
        background = add_picture_cover(slide, BACKGROUNDS / f"slide_{index:02d}.png", prs.slide_width, prs.slide_height)
        move_shape_to_index(slide, background, 2)
        if index != 1:
            add_title_veil(slide, prs.slide_width)

    prs.core_properties.title = "DriveClarify 研究进展汇报 V2 Image2"
    prs.core_properties.subject = "25-slide Image2 reconstruction with deterministic evidence overlays"
    prs.save(EDITABLE_PPTX)
    copy_support_files()
    write_image2_provenance()
    write_background_manifest()
    return EDITABLE_PPTX


def flatten(slides_dir: Path) -> Path:
    files = sorted(slides_dir.glob("slide-*.png"))
    if len(files) != 25:
        raise ValueError(f"Expected 25 rendered slide PNGs, found {len(files)} in {slides_dir}")
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    for path in files:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_picture_cover(slide, path, prs.slide_width, prs.slide_height)
    prs.core_properties.title = "DriveClarify 研究进展汇报 V2 Image2 全页图像版"
    prs.core_properties.subject = "Flattened deterministic export"
    prs.save(FLAT_PPTX)
    return FLAT_PPTX


def write_preservation_manifest(pptx_path: Path) -> Path:
    with zipfile.ZipFile(pptx_path) as archive:
        media_hashes: dict[str, list[str]] = {}
        for name in archive.namelist():
            if not name.startswith("ppt/media/"):
                continue
            digest = hashlib.sha256(archive.read(name)).hexdigest()
            media_hashes.setdefault(digest, []).append(name)

    records = []
    for label, path in REAL_ASSETS.items():
        digest = sha256(path)
        records.append(
            {
                "label": label,
                "source_path": str(path),
                "sha256": digest,
                "byte_identical_media_entries": media_hashes.get(digest, []),
                "preserved": bool(media_hashes.get(digest)),
            }
        )
    manifest = {
        "deck": str(pptx_path),
        "policy": "Image2 generated backgrounds only; exact text and evidence visuals remain deterministic foreground objects.",
        "assets": records,
    }
    output = OUT / "screenshot_preservation_manifest.json"
    output.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    if not all(record["preserved"] for record in records):
        missing = [record["label"] for record in records if not record["preserved"]]
        raise AssertionError(f"Evidence media not preserved byte-identically: {missing}")
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["compose", "flatten", "verify"])
    parser.add_argument("--slides-dir", type=Path, default=OUT / "rendered" / "slides")
    args = parser.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    if args.mode == "compose":
        path = compose()
        print(path)
    elif args.mode == "flatten":
        path = flatten(args.slides_dir)
        print(path)
    else:
        path = write_preservation_manifest(EDITABLE_PPTX)
        print(path)


if __name__ == "__main__":
    main()
